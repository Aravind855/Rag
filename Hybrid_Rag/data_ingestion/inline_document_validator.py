"""Inline Multi-Level Document Validator (Functional Approach).

Executes real-time structural (L1) and content/fidelity (L2) validation checks
during document ingestion inside Hybrid RAG, acting as an active gatekeeper before publishing.

Determines tri-state routing disposition:
- PUBLISH: Clean document exported to cleaned_documents/ for retrieval.
- QUARANTINE: Document has anomalies/warnings; isolated in quarantine/ for review.
- FAIL: Extraction failure or empty payload; routed to dead_letter/ queue.

Strictly procedural (zero classes). All paths resolved dynamically.
"""

import logging
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Tuple

# Optional parser imports for ground-truth structure inspection
try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import pymupdf
except ImportError:
    pymupdf = None

try:
    import pptx
except ImportError:
    pptx = None

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger("inline_document_validator")


# =====================================================================
# KNOWN ENTERPRISE DOMAIN GROUND TRUTH
# =====================================================================
CRITICAL_ENTITIES_BY_FILE = {
    "engineering/api_standards.md": ["NC-ENG-API-2026-08", "PROJECT_NOT_FOUND", "SCREAMING_SNAKE_CASE"],
    "operations/system_catalog.json": ["SYS-001", "SYS-002", "99.95%", "NC-OPS-SYS-2026-08"],
    "operations/vendor_management.docx": ["NC-OPS-VND-2026-08"],
    "security/access_control.md": ["NC-SEC-ACL-2026-08"],
    "security/incident_response.txt": ["NC-SEC-IR-2026-08"],
    "finance/expense_categories.xlsx": ["TRV-AIR", "TRV-HOT", "1000"],
    "finance/expense_policy.pdf": ["INR 1,000", "not reimbursable"],
    "hr/employee_handbook(1).pdf": ["INR 6,500", "INR 1,200", "INR 2,000"],
}


# =====================================================================
# LEVEL 1: DETERMINISTIC STRUCTURAL VALIDATION (100% COVERAGE)
# =====================================================================

def validate_l1_structural_integrity(
    clean_text: str,
    metadata: Dict[str, Any],
    raw_file_path: Path
) -> Dict[str, Any]:
    """Runs inexpensive deterministic structural integrity checks on the extracted document."""
    reasons: List[str] = []
    checks_passed = 0
    total_checks = 4

    # 1. Non-empty check on body content (excluding frontmatter header lines)
    body_text = clean_text
    if body_text.startswith("---"):
        parts = body_text.split("---", 2)
        if len(parts) >= 3:
            body_text = parts[2].strip()

    char_count = len(body_text)
    word_count = len(body_text.split())
    if char_count < 15 or word_count < 3:
        reasons.append(f"L1_FAIL: Empty or near-empty extracted body ({char_count} chars, {word_count} words).")
    else:
        checks_passed += 1

    # 2. Mandatory Metadata Schema Check
    required_meta = ["title", "category", "document_type", "source"]
    missing_meta = [k for k in required_meta if not metadata.get(k)]
    if missing_meta:
        reasons.append(f"L1_FAIL: Missing required metadata keys: {missing_meta}")
    else:
        checks_passed += 1

    # 3. Raw Source Readability & Size Sanity Check
    if not raw_file_path.exists():
        reasons.append(f"L1_FAIL: Raw source file missing at {raw_file_path}")
    elif raw_file_path.stat().st_size == 0:
        reasons.append("L1_FAIL: Raw source file is zero bytes.")
    else:
        checks_passed += 1

    # 4. Format-Specific Structural Unit Extraction
    ext = raw_file_path.suffix.lower()
    structural_unit_ok = True

    try:
        if ext == ".pdf" and pymupdf:
            doc = pymupdf.open(str(raw_file_path))
            pages = len(doc)
            doc.close()
            if pages == 0:
                structural_unit_ok = False
                reasons.append("L1_FAIL: PDF contains 0 pages.")

        elif ext == ".xlsx" and openpyxl:
            wb = openpyxl.load_workbook(str(raw_file_path), read_only=True)
            sheets = len(wb.sheetnames)
            wb.close()
            if sheets == 0:
                structural_unit_ok = False
                reasons.append("L1_FAIL: Excel workbook contains 0 sheets.")

        elif ext == ".pptx" and pptx:
            prs = pptx.Presentation(str(raw_file_path))
            if len(prs.slides) == 0:
                structural_unit_ok = False
                reasons.append("L1_FAIL: Presentation contains 0 slides.")
    except Exception as e:
        structural_unit_ok = False
        reasons.append(f"L1_FAIL: Failed structural parser check: {e}")

    if structural_unit_ok:
        checks_passed += 1

    passed = (len(reasons) == 0)
    score = (checks_passed / total_checks) * 100.0

    return {
        "passed": passed,
        "score": round(score, 2),
        "char_count": char_count,
        "word_count": word_count,
        "reasons": reasons,
    }


# =====================================================================
# LEVEL 2: CONTENT & FIDELITY VALIDATION
# =====================================================================

def validate_l2_content_fidelity(
    clean_text: str,
    raw_file_path: Path,
    source_rel_path: str
) -> Dict[str, Any]:
    """Runs deeper semantic and fidelity checks against the extracted markdown."""
    reasons: List[str] = []
    checks_passed = 0
    total_checks = 4

    clean_lower = clean_text.lower()

    # 1. Critical Domain Entities
    expected_entities = CRITICAL_ENTITIES_BY_FILE.get(source_rel_path, [])
    if expected_entities:
        missing_entities = [e for e in expected_entities if e.lower() not in clean_lower]
        if missing_entities:
            reasons.append(f"L2_WARN: Missing critical domain entities: {missing_entities}")
        else:
            checks_passed += 1
    else:
        # Generic identifier check
        has_identifiers = bool(re.search(r"\b[A-Z0-9_-]{4,}\b", clean_text))
        if has_identifiers:
            checks_passed += 1
        else:
            reasons.append("L2_WARN: Document lacks standard capitalized identifiers or codes.")

    # 2. Markdown Table Structural Integrity (Per Table Block)
    lines = clean_text.splitlines()
    table_blocks: List[List[str]] = []
    current_block: List[str] = []

    for line in lines:
        s = line.strip()
        if s.startswith("|") and s.endswith("|"):
            current_block.append(s)
        else:
            if current_block:
                table_blocks.append(current_block)
                current_block = []
    if current_block:
        table_blocks.append(current_block)

    tables_misaligned = 0
    for block in table_blocks:
        data_rows = [r for r in block if not re.match(r"^\|(\s*:?-+:?\s*\|)+$", r)]
        if len(data_rows) >= 2:
            col_counts = [len(r.split("|")[1:-1]) for r in data_rows]
            if len(set(col_counts)) > 1:
                tables_misaligned += 1

    if tables_misaligned == 0:
        checks_passed += 1
    else:
        reasons.append(f"L2_WARN: Markdown table column misalignment detected in {tables_misaligned} tables.")

    # 3. Heading Tree Integrity (No severe level skips like H1 directly to H4)
    headings = [len(m.group(1)) for m in re.finditer(r"^(#{1,6})\s+", clean_text, flags=re.MULTILINE)]
    has_level_skip = False
    for i in range(len(headings) - 1):
        if headings[i + 1] > headings[i] + 2:
            has_level_skip = True
            break

    if has_level_skip:
        reasons.append("L2_WARN: Skipped heading levels detected (e.g. H1 followed immediately by H4).")
    else:
        checks_passed += 1

    # 4. Content Volume Sanity (Drop Ratio)
    raw_size = raw_file_path.stat().st_size if raw_file_path.exists() else 0
    clean_size = len(clean_text)

    ext = raw_file_path.suffix.lower()
    if ext in [".md", ".txt", ".json"]:
        if raw_size > 0 and (clean_size / raw_size) < 0.60:
            reasons.append(f"L2_WARN: Text drop anomaly: Clean text length ({clean_size}) is < 60% of raw text ({raw_size}).")
        else:
            checks_passed += 1
    else:
        if clean_size >= 50:
            checks_passed += 1
        else:
            reasons.append(f"L2_WARN: Extracted text volume suspiciously low ({clean_size} chars) for {ext} file.")

    score = (checks_passed / total_checks) * 100.0

    return {
        "score": round(score, 2),
        "checks_passed": checks_passed,
        "total_checks": total_checks,
        "reasons": reasons,
    }


# =====================================================================
# TRI-STATE DISPOSITION ENGINE
# =====================================================================

def evaluate_document_for_ingestion(
    clean_text: str,
    metadata: Dict[str, Any],
    raw_file_path: Path,
    source_rel_path: str
) -> Tuple[str, float, List[str], Dict[str, Any]]:
    """Evaluates document through L1 & L2 validation gates.

    Returns:
        Tuple of (disposition, composite_score, reasons_list, details_dict)
    """
    l1 = validate_l1_structural_integrity(clean_text, metadata, raw_file_path)
    l2 = validate_l2_content_fidelity(clean_text, raw_file_path, source_rel_path)

    all_reasons = l1["reasons"] + l2["reasons"]
    composite_score = round((l1["score"] * 0.50) + (l2["score"] * 0.50), 2)

    if not l1["passed"]:
        disposition = "FAIL"
    elif composite_score >= 90.0 and not any(r.startswith("L1_") for r in all_reasons):
        disposition = "PUBLISH"
    elif composite_score >= 70.0:
        disposition = "QUARANTINE"
    else:
        disposition = "FAIL"

    details = {
        "l1_structural": l1,
        "l2_fidelity": l2,
        "composite_score": composite_score,
    }

    return disposition, composite_score, all_reasons, details
